# ============================================================
# journal.py — Trade Journal Logger
# ============================================================

import os
import pandas as pd
from datetime import datetime
from openpyxl import load_workbook

JOURNAL_FILE = "Alpaca_Trading_Journal.xlsx"
SHEET_NAME   = "Trade Journal"


def _read_journal():
    """Read the journal Excel file, automatically skipping the title row if present."""
    df = pd.read_excel(JOURNAL_FILE, sheet_name=SHEET_NAME, header=0)
    # If first column header is not 'Date', there's a title row — skip it
    if str(df.columns[0]).strip() != "Date":
        df = pd.read_excel(JOURNAL_FILE, sheet_name=SHEET_NAME, header=1)
    # Drop completely empty rows
    df = df.dropna(how="all")
    return df


def log_trade(symbol, side, qty, entry_price, stop_loss,
              take_profit=None, order_type="Market",
              reason="", notes=""):
    """
    Log a new trade entry to the Excel journal.
    Call this when you OPEN a trade.
    """
    row = {
        "Date":            datetime.now().strftime("%Y-%m-%d"),
        "Symbol":          symbol,
        "Side":            side.upper(),
        "Qty":             qty,
        "Entry Price ($)": round(entry_price, 2),
        "Exit Price ($)":  "",
        "Stop Loss ($)":   round(stop_loss, 2),
        "Take Profit ($)": round(take_profit, 2) if take_profit else "",
        "P&L ($)":         "",
        "P&L (%)":         "",
        "Status":          "Open",
        "Order Type":      order_type,
        "Reason / Strategy": reason,
        "Notes / Lessons": notes,
        "Exit Date":       "",
    }

    _append_row(row)
    print(f"[JOURNAL] Trade logged: {side.upper()} {qty} {symbol} @ ${entry_price:.2f}")


def close_trade(symbol, exit_price, notes="", qty=None):
    """
    Update the most recent open trade for a symbol with exit price.
    Call this when you CLOSE a trade.

    `qty` = the shares actually sold in THIS close. Pass it whenever the caller
    knows the real remaining size (the agent always does). If omitted, the row's
    recorded Qty is used. Passing it prevents double-counting: a position that was
    scaled out earlier still has its ORIGINAL Qty on the Open row, so computing P&L
    on that would count the already-sold shares a second time.
    """
    if not os.path.exists(JOURNAL_FILE):
        print(f"[JOURNAL] File not found: {JOURNAL_FILE}")
        return

    try:
        df = _read_journal()

        # Find the most recent open trade for this symbol
        mask = (df["Symbol"] == symbol) & (df["Status"] == "Open")
        if not mask.any():
            print(f"[JOURNAL] No open trade found for {symbol}")
            return

        idx = df[mask].index[-1]
        entry_price = float(df.at[idx, "Entry Price ($)"])
        side        = df.at[idx, "Side"]
        close_qty   = int(qty) if qty is not None else float(df.at[idx, "Qty"])

        # Calculate P&L on the shares actually sold in this close
        if side == "BUY":
            pnl     = (exit_price - entry_price) * close_qty
            pnl_pct = (exit_price - entry_price) / entry_price * 100
        else:
            pnl     = (entry_price - exit_price) * close_qty
            pnl_pct = (entry_price - exit_price) / entry_price * 100

        df.at[idx, "Qty"]            = close_qty
        df.at[idx, "Exit Price ($)"] = round(exit_price, 2)
        df.at[idx, "P&L ($)"]        = round(pnl, 2)
        df.at[idx, "P&L (%)"]        = round(pnl_pct, 2)
        df.at[idx, "Status"]          = "Closed"
        df.at[idx, "Exit Date"]      = datetime.now().strftime("%Y-%m-%d")
        if notes:
            df.at[idx, "Notes / Lessons"] = notes

        # Write back
        with pd.ExcelWriter(JOURNAL_FILE, engine="openpyxl",
                            mode="a", if_sheet_exists="replace") as writer:
            df.to_excel(writer, sheet_name=SHEET_NAME, index=False)

        sign = "+" if pnl >= 0 else ""
        print(f"[JOURNAL] Trade closed: {symbol} @ ${exit_price:.2f} | "
              f"P&L: {sign}${pnl:.2f} ({sign}{pnl_pct:.1f}%)")

    except Exception as e:
        print(f"[JOURNAL] Error closing trade: {e}")


def _append_row(row):
    """Append a row to the journal Excel file."""
    if not os.path.exists(JOURNAL_FILE):
        # Create new file with headers
        df = pd.DataFrame([row])
        df.to_excel(JOURNAL_FILE, sheet_name=SHEET_NAME, index=False)
        print(f"[JOURNAL] Created new journal file: {JOURNAL_FILE}")
        return

    try:
        wb = load_workbook(JOURNAL_FILE)
        ws = wb[SHEET_NAME]
        # Find the last row with data and append after it
        last_row = ws.max_row + 1
        for col_idx, val in enumerate(row.values(), 1):
            ws.cell(row=last_row, column=col_idx, value=val)
        wb.save(JOURNAL_FILE)
    except Exception as e:
        print(f"[JOURNAL] Error writing to journal: {e}")


def log_partial_sell(symbol, sell_qty, entry_price, exit_price, reason=""):
    """Log a partial sell (profit-take) as a new row in the journal."""
    pnl     = (exit_price - entry_price) * sell_qty
    pnl_pct = (exit_price - entry_price) / entry_price * 100
    row = {
        "Date":              datetime.now().strftime("%Y-%m-%d"),
        "Symbol":            symbol,
        "Side":              "SELL",
        "Qty":               sell_qty,
        "Entry Price ($)":   round(entry_price, 2),
        "Exit Price ($)":    round(exit_price, 2),
        "Stop Loss ($)":     "",
        "Take Profit ($)":   "",
        "P&L ($)":           round(pnl, 2),
        "P&L (%)":           round(pnl_pct, 2),
        "Status":            "Partial Close",
        "Order Type":        "Market",
        "Reason / Strategy": reason,
        "Notes / Lessons":   "",
        "Exit Date":         datetime.now().strftime("%Y-%m-%d"),
    }
    _append_row(row)
    sign = "+" if pnl >= 0 else ""
    print(f"[JOURNAL] Partial sell logged: SELL {sell_qty} {symbol} @ ${exit_price:.2f} | "
          f"P&L: {sign}${pnl:.2f} ({sign}{pnl_pct:.1f}%)")


def print_summary():
    """Print a quick summary of your trading performance."""
    if not os.path.exists(JOURNAL_FILE):
        print("[JOURNAL] No journal file found yet.")
        return

    try:
        df = _read_journal()
        closed = df[df["Status"] == "Closed"].copy()

        if closed.empty:
            print("[JOURNAL] No closed trades yet.")
            return

        closed["P&L ($)"] = pd.to_numeric(closed["P&L ($)"], errors="coerce")
        total_trades  = len(closed)
        winners       = len(closed[closed["P&L ($)"] > 0])
        losers        = len(closed[closed["P&L ($)"] < 0])
        win_rate      = (winners / total_trades * 100) if total_trades > 0 else 0
        total_pnl     = closed["P&L ($)"].sum()
        best_trade    = closed["P&L ($)"].max()
        worst_trade   = closed["P&L ($)"].min()

        print("\n" + "="*50)
        print("  TRADING JOURNAL SUMMARY")
        print("="*50)
        print(f"  Total closed trades : {total_trades}")
        print(f"  Winners             : {winners}")
        print(f"  Losers              : {losers}")
        print(f"  Win rate            : {win_rate:.1f}%")
        print(f"  Total P&L           : ${total_pnl:+.2f}")
        print(f"  Best trade          : ${best_trade:+.2f}")
        print(f"  Worst trade         : ${worst_trade:+.2f}")
        print("="*50 + "\n")

    except Exception as e:
        print(f"[JOURNAL] Error reading journal: {e}")
